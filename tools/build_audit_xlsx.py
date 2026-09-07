#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""生成/刷新 循环点审计.xlsx：逐曲审计 + 循环点权威比对 + 按作品汇总"""
import csv
import gzip
import json
import os
import statistics as st

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BASE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
idx = json.loads(gzip.open(os.path.join(BASE, 'src', 'Resources', 'tracks.tf.json.gz'), 'rb').read().decode('utf-8'))
code = {g['id']: g['code'] for g in idx['games']}
gname = {g['id']: g['name'] for g in idx['games']}


def I(v):
    try:
        return int(float(v))
    except Exception:
        return None


def level(jump, k, lendiff):
    if jump is None:
        return "—"
    pct = jump / 32768 * 100
    if pct >= 40 or abs(k or 0) >= 1500 or abs(lendiff or 0) >= 3000:
        return "严重"
    if pct >= 20 or abs(k or 0) >= 600 or abs(lendiff or 0) >= 1000:
        return "中等"
    if pct >= 8 or abs(k or 0) >= 200 or abs(lendiff or 0) >= 200:
        return "轻微"
    return "正常"


wb = openpyxl.Workbook()

# ── 页1：逐曲审计 ──
ws = wb.active
ws.title = "逐曲审计"
heads = ["作品代号", "作品名", "条目文件", "曲名", "循环起点(秒)", "循环终点(秒)", "索引时长(秒)",
         "解码样本数", "与真实长度差(样本)", "与真实长度差(毫秒)", "折返跳变(幅度)", "折返跳变(满刻度%)",
         "最优偏移(样本)", "最优偏移(毫秒)", "校正后残差", "问题等级"]
ws.append(heads)
for c in ws[1]:
    c.font = Font(bold=True)
    c.alignment = Alignment(horizontal="center", vertical="center")
    c.fill = PatternFill("solid", fgColor="DDEBF7")
ws.freeze_panes = "A2"
fills = {"严重": "F8CBAD", "中等": "FFE699", "轻微": "FFF2CC", "正常": "E2EFDA", "—": "F2F2F2"}

rows = list(csv.DictReader(open(os.path.join(BASE, 'docs/loop_audit.csv'), encoding='utf-8-sig')))
for r in rows:
    gid = r['game']
    jump, k, ld = I(r['jump']), I(r['best_k']), I(r['len_diff'])
    ws.append([code.get(gid, gid), gname.get(gid, ""), r['file'], r['title'],
               float(r['ls']) if r['ls'] else None, float(r['le']) if r['le'] else None,
               float(r['d']) if r['d'] else None, I(r['full_samples']), ld,
               round(ld / 44.1, 2) if ld is not None else None, jump,
               round(jump / 32768 * 100, 1) if jump is not None else None, k,
               round(k / 44.1, 2) if k is not None else None, I(r['residual']),
               level(jump, k, ld)])
    ws.cell(row=ws.max_row, column=16).fill = PatternFill("solid", fgColor=fills[level(jump, k, ld)])
for i, w in enumerate([10, 42, 18, 34, 12, 12, 12, 13, 18, 18, 14, 16, 14, 14, 12, 10], 1):
    ws.column_dimensions[get_column_letter(i)].width = w
ws.auto_filter.ref = ws.dimensions

# ── 页2：循环点权威比对 ──
ws2 = wb.create_sheet("循环点权威比对")
h2 = ["作品代号", "条目文件", "曲名", "游戏自带起点(样本)", "游戏自带终点(样本)",
      "索引起点(样本)", "索引终点(样本)", "Δ起点", "Δ终点", "Δ起点(ms)", "Δ终点(ms)"]
ws2.append(h2)
for c in ws2[1]:
    c.font = Font(bold=True)
    c.fill = PatternFill("solid", fgColor="DDEBF7")
    c.alignment = Alignment(horizontal="center")
drows = list(csv.DictReader(open(os.path.join(BASE, 'docs/loop_delta.csv'), encoding='utf-8-sig')))
for r in drows[1:]:
    ds, de = I(r['d_start(样本)']), I(r['d_end(样本)'])
    ws2.append([code.get(r['game'], r['game']), r['file'], r['title'], I(r['auth_start']), I(r['auth_end']),
                I(r['idx_start']), I(r['idx_end']), ds, de,
                round(ds / 44.1, 3) if ds is not None else None,
                round(de / 44.1, 3) if de is not None else None])
    if ds is not None and abs(ds) > 30:
        ws2.cell(row=ws2.max_row, column=8).fill = PatternFill("solid", fgColor="F8CBAD")
for i, w in enumerate([10, 18, 34, 18, 18, 14, 14, 10, 10, 12, 12], 1):
    ws2.column_dimensions[get_column_letter(i)].width = w
ws2.freeze_panes = "A2"

# ── 页3：按作品汇总 ──
ws3 = wb.create_sheet("按作品汇总")
h3 = ["作品代号", "作品名", "曲目数", "循环曲", "折返跳变中位", "折返跳变最大",
      "长度差中位(样本)", "长度差最大(样本)", "严重/中等/轻微 曲数"]
ws3.append(h3)
for c in ws3[1]:
    c.font = Font(bold=True)
    c.fill = PatternFill("solid", fgColor="DDEBF7")
    c.alignment = Alignment(horizontal="center")
for gid in ['th075', 'th105', 'th123', 'th135', 'th145', 'th155', 'th175']:
    rs = [r for r in rows if r['game'] == gid]
    lp = [r for r in rs if I(r['jump']) is not None]
    jm = [I(r['jump']) for r in lp]
    ld = [I(r['len_diff']) for r in rs if I(r['len_diff']) is not None]
    lvs = [level(I(r['jump']), I(r['best_k']), I(r['len_diff'])) for r in rs]
    ws3.append([code[gid], gname[gid], len(rs), len(lp),
                int(st.median(jm)) if jm else 0, max(jm) if jm else 0,
                int(st.median(ld)) if ld else 0, max(ld, key=abs) if ld else 0,
                f"{lvs.count('严重')}/{lvs.count('中等')}/{lvs.count('轻微')}"])
for i, w in enumerate([10, 42, 9, 9, 14, 14, 18, 18, 20], 1):
    ws3.column_dimensions[get_column_letter(i)].width = w

out = os.path.join(BASE, 'docs/循环点审计.xlsx')
wb.save(out)
print('已生成', out, f'（逐曲 {len(rows)} 行 / 权威比对 {len(drows) - 1} 行）')
